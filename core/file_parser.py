import io
from pypdf import PdfReader
from docx import Document

from config import MAX_TOTAL_CHARS


class UnsupportedFileError(Exception):
    pass


class EmptyDocumentError(Exception):
    pass

def _extract_from_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages_text = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        pages_text.append(page_text)
    return "\n".join(pages_text)


def _extract_from_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)

    return "\n".join(paragraphs)


def extract_text(uploaded_file) -> str:
    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.read()

    if filename.endswith(".pdf"):
        text = _extract_from_pdf(file_bytes)
    elif filename.endswith(".docx"):
        text = _extract_from_docx(file_bytes)
    else:
        raise UnsupportedFileError(
            "Unsupported file type. Please upload a .pdf or .docx file."
        )

    text = text.strip()
    if not text:
        raise EmptyDocumentError(
            "No readable text was found in this document. "
            "It may be a scanned/image-only file."
        )

    if len(text) > MAX_TOTAL_CHARS:
        text = text[:MAX_TOTAL_CHARS]

    return text


def chunk_text(text: str, max_chars: int) -> list[str]:
    if len(text) <= max_chars:
        return [text]

    paragraphs = text.split("\n")
    chunks = []
    current = ""

    for para in paragraphs:
        if len(current) + len(para) + 1 <= max_chars:
            current += para + "\n"
        else:
            if current.strip():
                chunks.append(current.strip())
            if len(para) > max_chars:
                for i in range(0, len(para), max_chars):
                    chunks.append(para[i:i + max_chars])
                current = ""
            else:
                current = para + "\n"

    if current.strip():
        chunks.append(current.strip())

    return chunks