import io

from docx import Document
from docx.shared import RGBColor
from fpdf import FPDF
from fpdf.enums import XPos, YPos

from core.question_generator import QuizResult


def build_docx(quiz: QuizResult, title: str = "Generated Quiz") -> bytes:
    doc = Document()
    doc.add_heading(title, level=0)

    if quiz.mcqs:
        doc.add_heading("Multiple Choice Questions", level=1)
        for i, q in enumerate(quiz.mcqs, start=1):
            p = doc.add_paragraph()
            run = p.add_run(f"{i}. {q.question}")
            run.bold = True

            for j, opt in enumerate(q.options):
                label = chr(65 + j)
                doc.add_paragraph(f"   {label}. {opt}")

            ans_p = doc.add_paragraph()
            ans_run = ans_p.add_run(f"Correct Answer: {q.correct_answer}")
            ans_run.font.color.rgb = RGBColor(0, 128, 0)
            ans_run.bold = True

            if q.explanation:
                exp_p = doc.add_paragraph()
                exp_run = exp_p.add_run(f"Explanation: {q.explanation}")
                exp_run.italic = True
            doc.add_paragraph()

    if quiz.short_answers:
        doc.add_heading("Short Answer Questions", level=1)
        for i, q in enumerate(quiz.short_answers, start=1):
            p = doc.add_paragraph()
            run = p.add_run(f"{i}. {q.question}")
            run.bold = True

            ans_p = doc.add_paragraph()
            ans_run = ans_p.add_run(f"Ideal Answer: {q.ideal_answer}")
            ans_run.font.color.rgb = RGBColor(0, 128, 0)
            doc.add_paragraph()

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _clean_for_pdf(text: str) -> str:
    return text.encode("latin-1", errors="replace").decode("latin-1")


def build_pdf(quiz: QuizResult, title: str = "Generated Quiz") -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)

    pdf.set_font("Helvetica", "B", 18)
    pdf.multi_cell(0, 10, _clean_for_pdf(title), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(4)

    if quiz.mcqs:
        pdf.set_font("Helvetica", "B", 14)
        pdf.multi_cell(0, 10, "Multiple Choice Questions", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(2)

        for i, q in enumerate(quiz.mcqs, start=1):
            pdf.set_font("Helvetica", "B", 11)
            pdf.multi_cell(0, 7, _clean_for_pdf(f"{i}. {q.question}"), new_x=XPos.LMARGIN, new_y=YPos.NEXT)

            pdf.set_font("Helvetica", "", 11)
            for j, opt in enumerate(q.options):
                label = chr(65 + j)
                pdf.multi_cell(0, 7, _clean_for_pdf(f"   {label}. {opt}"), new_x=XPos.LMARGIN, new_y=YPos.NEXT)

            pdf.set_font("Helvetica", "B", 11)
            pdf.set_text_color(0, 128, 0)
            pdf.multi_cell(0, 7, _clean_for_pdf(f"Correct Answer: {q.correct_answer}"), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.set_text_color(0, 0, 0)

            if q.explanation:
                pdf.set_font("Helvetica", "I", 10)
                pdf.multi_cell(0, 7, _clean_for_pdf(f"Explanation: {q.explanation}"), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.ln(3)

    if quiz.short_answers:
        pdf.set_font("Helvetica", "B", 14)
        pdf.multi_cell(0, 10, "Short Answer Questions", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(2)

        for i, q in enumerate(quiz.short_answers, start=1):
            pdf.set_font("Helvetica", "B", 11)
            pdf.multi_cell(0, 7, _clean_for_pdf(f"{i}. {q.question}"), new_x=XPos.LMARGIN, new_y=YPos.NEXT)

            pdf.set_font("Helvetica", "", 11)
            pdf.set_text_color(0, 128, 0)
            pdf.multi_cell(0, 7, _clean_for_pdf(f"Ideal Answer: {q.ideal_answer}"), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.set_text_color(0, 0, 0)
            pdf.ln(3)

    return bytes(pdf.output())